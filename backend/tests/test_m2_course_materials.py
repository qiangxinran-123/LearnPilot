from __future__ import annotations

import io
import json
import os
import subprocess
import sys
from pathlib import Path
from urllib.error import HTTPError

import pymupdf
import pytest
from docx import Document
from pydantic import ValidationError
from sqlalchemy import create_engine, inspect, select
from sqlalchemy.orm import Session, sessionmaker

from backend.app.ai.client import OpenAICompatibleLLMClient
from backend.app.ai.errors import AIServiceError
from backend.app.ai.models import SummaryOutput
from backend.app.core.settings import Settings
from backend.app.db.models import Course, Material, MaterialContent, MaterialSummary, StudyCard, User
from backend.app.documents import ParseStatus, parse_document
from backend.app.repositories import (
    CourseRepository,
    MaterialContentRepository,
    MaterialRepository,
    MaterialSummaryRepository,
    StudyCardRepository,
    UserRepository,
)
from backend.app.schemas.course import CourseCreate, CourseUpdate
from backend.app.services.course_service import CourseService, ResourceNotFoundError
from backend.app.services.material_ai_service import MaterialAIService
from backend.app.storage import LocalFileStorage
from backend.scripts.import_legacy_json import import_legacy_data


def migrate_database(tmp_path: Path, revision: str = "head") -> str:
    database_url = f"sqlite+pysqlite:///{tmp_path / 'm2.sqlite3'}"
    environment = os.environ.copy()
    environment["DATABASE_URL"] = database_url
    subprocess.run(
        [sys.executable, "-m", "alembic", "upgrade", revision],
        env=environment,
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    return database_url


def make_settings(database_url: str, upload_root: Path, *, llm_mode: str = "mock") -> Settings:
    return Settings(
        app_env="test",
        database_url=database_url,
        upload_root=upload_root,
        max_upload_bytes=1024 * 1024,
        llm_mode=llm_mode,
        llm_api_key="",
        llm_base_url="",
        llm_model="",
        llm_context_chars=2_000,
        llm_chunk_chars=400,
        llm_chunk_overlap_chars=80,
    )


def make_course_service(session: Session, settings: Settings) -> CourseService:
    return CourseService(
        session,
        CourseRepository(session),
        MaterialRepository(session),
        MaterialContentRepository(session),
        UserRepository(session),
        LocalFileStorage(settings.upload_root),
        settings,
    )


def make_ai_service(session: Session, settings: Settings) -> MaterialAIService:
    return MaterialAIService(
        session,
        MaterialRepository(session),
        MaterialContentRepository(session),
        MaterialSummaryRepository(session),
        StudyCardRepository(session),
        UserRepository(session),
        settings,
    )


def make_pdf_bytes(text: str) -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    payload = document.tobytes()
    document.close()
    return payload


def make_docx_bytes(paragraph: str, table_value: str) -> bytes:
    document = Document()
    document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = table_value
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_m2_migration_upgrades_m1_schema_to_artifact_schema(tmp_path: Path) -> None:
    database_url = migrate_database(tmp_path, "0001_m1a_foundation")
    engine = create_engine(database_url)
    assert "material_contents" not in inspect(engine).get_table_names()
    engine.dispose()

    migrate_database(tmp_path)
    engine = create_engine(database_url)
    tables = set(inspect(engine).get_table_names())
    assert {"material_contents", "material_summaries", "study_cards"}.issubset(tables)
    assert inspect(engine).get_foreign_keys("material_contents")[0]["referred_table"] == "materials"
    engine.dispose()


def test_document_parsers_extract_pdf_docx_and_report_failures() -> None:
    pdf = parse_document(io.BytesIO(make_pdf_bytes("Cache locality reduces latency.")), ".pdf")
    docx = parse_document(io.BytesIO(make_docx_bytes("Pipeline stages", "hazard table")), ".docx")
    corrupt_pdf = parse_document(io.BytesIO(b"%PDF-1.4 not a complete PDF"), ".pdf")
    invalid_text = parse_document(io.BytesIO(bytes([255, 254])), ".txt")
    empty_markdown = parse_document(io.BytesIO(b"\n\n"), ".md")

    assert pdf.status == ParseStatus.SUCCESS
    assert pdf.page_count == 1
    assert "Cache locality" in pdf.text
    assert pdf.segments[0].source_reference == "PDF 第 1 页"
    assert docx.status == ParseStatus.SUCCESS
    assert "Pipeline stages" in docx.text
    assert "hazard table" in docx.text
    assert corrupt_pdf.status == ParseStatus.PARSE_FAILED
    assert corrupt_pdf.error_code == "PDF_PARSE_ERROR"
    assert invalid_text.status == ParseStatus.DECODE_ERROR
    assert empty_markdown.status == ParseStatus.EMPTY_CONTENT


def test_course_crud_validates_input_and_preserves_owner_boundary(tmp_path: Path) -> None:
    database_url = migrate_database(tmp_path)
    engine = create_engine(database_url)
    session = sessionmaker(bind=engine, expire_on_commit=False)()
    settings = make_settings(database_url, tmp_path / "uploads")
    service = make_course_service(session, settings)

    with pytest.raises(ValidationError):
        CourseCreate(name="  ")
    course = service.create_course(
        CourseCreate(name="Computer Systems", description="Memory and processors", category="CS")
    )
    course_id = course.id
    updated = service.update_course(course_id, CourseUpdate(learning_goal="Understand cache behavior"))
    assert updated.learning_goal == "Understand cache behavior"

    with session.begin():
        other_owner = UserRepository(session).add(
            User(email="other-owner@example.test", display_name="Other", is_active=True)
        )
        private_course = CourseRepository(session).add(Course(owner_id=other_owner.id, name="Private", description=""))
    with pytest.raises(ResourceNotFoundError):
        service.get_course(private_course.id)

    service.delete_course(course_id)
    with pytest.raises(ResourceNotFoundError):
        service.get_course(course_id)
    session.close()
    engine.dispose()


def test_upload_persists_parsed_content_and_course_delete_removes_storage(tmp_path: Path) -> None:
    database_url = migrate_database(tmp_path)
    engine = create_engine(database_url)
    session_factory = sessionmaker(bind=engine, expire_on_commit=False)
    settings = make_settings(database_url, tmp_path / "uploads")

    with session_factory() as session:
        service = make_course_service(session, settings)
        course = service.create_course(CourseCreate(name="Architecture"))
        material = service.create_uploaded_material(
            course.id,
            "cache-notes.txt",
            "text/plain",
            io.BytesIO(b"Cache locality improves average memory access time."),
        )
        storage_key = material.storage_key
        assert material.extraction_status == "SUCCESS"
        assert storage_key is not None
        assert service.storage.exists(storage_key)
        content = service.get_content(material.id)
        assert content.content_hash
        assert "Cache locality" in content.plain_text
        material_id = material.id

    with session_factory() as session:
        service = make_course_service(session, settings)
        assert "Cache locality" in service.get_content(material_id).plain_text
        service.delete_course(course.id)
        assert not service.storage.exists(storage_key)
        assert MaterialRepository(session).get_by_id(material_id) is None

    engine.dispose()


def test_pdf_docx_uploads_persist_parser_provenance_and_no_path(tmp_path: Path) -> None:
    database_url = migrate_database(tmp_path)
    engine = create_engine(database_url)
    session = sessionmaker(bind=engine, expire_on_commit=False)()
    settings = make_settings(database_url, tmp_path / "uploads")
    service = make_course_service(session, settings)
    course = service.create_course(CourseCreate(name="Documents"))

    pdf = service.create_uploaded_material(
        course.id,
        "lecture.pdf",
        "application/pdf",
        io.BytesIO(make_pdf_bytes("Virtual memory maps pages.")),
    )
    docx = service.create_uploaded_material(
        course.id,
        "notes.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        io.BytesIO(make_docx_bytes("Instruction pipeline", "stall")),
    )
    assert pdf.extraction_status == "SUCCESS"
    assert docx.extraction_status == "SUCCESS"
    assert service.get_content(pdf.id).page_count == 1
    assert service.get_content(docx.id).parser_name == "python-docx"
    session.close()
    engine.dispose()


def test_ai_artifacts_persist_and_material_qa_stays_grounded(tmp_path: Path) -> None:
    database_url = migrate_database(tmp_path)
    engine = create_engine(database_url)
    session_factory = sessionmaker(bind=engine, expire_on_commit=False)
    settings = make_settings(database_url, tmp_path / "uploads", llm_mode="mock")

    with session_factory() as session:
        course_service = make_course_service(session, settings)
        course = course_service.create_course(CourseCreate(name="Memory"))
        material = course_service.create_uploaded_material(
            course.id,
            "cache.txt",
            "text/plain",
            io.BytesIO(b"Cache locality improves average memory access time and reduces latency."),
        )
        ai_service = make_ai_service(session, settings)
        summary = ai_service.generate_summary(material.id)
        cards = ai_service.generate_cards(material.id)
        answer = ai_service.answer_question(material.id, "What does cache locality improve?")
        material_id = material.id
        assert summary.model_provider == "mock"
        assert summary.prompt_version == "summary_v1"
        assert len(cards) == 5
        assert all(card.prompt_version == "cards_v3" for card in cards)
        assert answer.grounded
        assert answer.references
        assert answer.prompt_version == "material_qa_v1"

    with session_factory() as session:
        assert MaterialSummaryRepository(session).get_for_material(material_id) is not None
        assert len(StudyCardRepository(session).list_for_material(material_id)) == 5
        assert MaterialContentRepository(session).get_for_material(material_id) is not None
    engine.dispose()


def test_unconfigured_llm_never_silently_uses_mock(tmp_path: Path) -> None:
    database_url = migrate_database(tmp_path)
    engine = create_engine(database_url)
    session = sessionmaker(bind=engine, expire_on_commit=False)()
    mock_settings = make_settings(database_url, tmp_path / "uploads", llm_mode="mock")
    course_service = make_course_service(session, mock_settings)
    course = course_service.create_course(CourseCreate(name="AI configuration"))
    material = course_service.create_uploaded_material(
        course.id,
        "source.txt",
        "text/plain",
        io.BytesIO(b"A cache stores recently used data."),
    )
    provider_settings = make_settings(database_url, tmp_path / "uploads", llm_mode="provider")
    with pytest.raises(AIServiceError) as error:
        make_ai_service(session, provider_settings).generate_summary(material.id)
    assert error.value.code == "AI_NOT_CONFIGURED"
    session.close()
    engine.dispose()


def test_openai_compatible_adapter_normalizes_timeout_rate_and_invalid_output(monkeypatch) -> None:
    settings = Settings(
        llm_mode="provider",
        llm_api_key="test-key",
        llm_base_url="https://example.test/v1",
        llm_model="test-model",
        llm_max_attempts=1,
    )
    client = OpenAICompatibleLLMClient(settings)

    def timeout(*args, **kwargs):
        raise TimeoutError

    monkeypatch.setattr("backend.app.ai.client.urlopen", timeout)
    with pytest.raises(AIServiceError) as timeout_error:
        client.generate_structured(operation="test", messages=[], schema=SummaryOutput)
    assert timeout_error.value.code == "AI_TIMEOUT"

    def rate_limited(*args, **kwargs):
        raise HTTPError("https://example.test", 429, "rate limited", {}, None)

    monkeypatch.setattr("backend.app.ai.client.urlopen", rate_limited)
    with pytest.raises(AIServiceError) as rate_error:
        client.generate_structured(operation="test", messages=[], schema=SummaryOutput)
    assert rate_error.value.code == "AI_RATE_LIMITED"

    class FakeResponse:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def read(self) -> bytes:
            return json.dumps({"choices": [{"message": {"content": "{}"}}]}).encode()

    monkeypatch.setattr("backend.app.ai.client.urlopen", lambda *args, **kwargs: FakeResponse())
    with pytest.raises(AIServiceError) as invalid_error:
        client.generate_structured(operation="test", messages=[], schema=SummaryOutput)
    assert invalid_error.value.code == "AI_INVALID_OUTPUT"


def test_populated_m1_upgrade_backfills_full_text_and_marks_partial_legacy_data(tmp_path: Path) -> None:
    database_url = migrate_database(tmp_path, "0001_m1a_foundation")
    engine = create_engine(database_url)
    session_factory = sessionmaker(bind=engine)
    with session_factory() as session:
        with session.begin():
            owner = User(email="upgrade-owner@example.test", display_name="Upgrade", is_active=True)
            session.add(owner)
            session.flush()
            course = Course(owner_id=owner.id, name="Legacy systems", description="")
            session.add(course)
            session.flush()
            session.add_all(
                [
                    Material(
                        id=801,
                        course_id=course.id,
                        title="Full M1 text",
                        material_type="txt",
                        summary="",
                        storage_status="STORED",
                        extraction_status="EXTRACTED",
                        text_extracted=True,
                        text_preview="Memory hierarchy",
                        text_length=len("Memory hierarchy"),
                        extracted_text="Memory hierarchy",
                        ai_summary="Old mock overview",
                        summary_generated=True,
                        summary_method="mock",
                        cards=[
                            {"question": "What is memory hierarchy?", "answer": "An old mock answer."},
                        ],
                        cards_generated=True,
                        cards_method="mock",
                    ),
                    Material(
                        id=802,
                        course_id=course.id,
                        title="Truncated M1 text",
                        material_type="txt",
                        summary="",
                        storage_status="STORED",
                        extraction_status="EXTRACTED",
                        text_extracted=True,
                        text_preview="partial",
                        text_length=25_000,
                        extracted_text="partial",
                    ),
                    Material(
                        id=803,
                        course_id=course.id,
                        title="Preview only",
                        material_type="txt",
                        summary="",
                        storage_status="LEGACY_FILE_UNMIGRATED",
                        extraction_status="EXTRACTED",
                        text_extracted=True,
                        text_preview="Preview is not the full document",
                        text_length=100,
                    ),
                ]
            )
    engine.dispose()

    migrate_database(tmp_path)
    engine = create_engine(database_url)
    with sessionmaker(bind=engine)() as session:
        full = session.get(Material, 801)
        full_content = session.scalar(select(MaterialContent).where(MaterialContent.material_id == 801))
        legacy_summary = session.scalar(select(MaterialSummary).where(MaterialSummary.material_id == 801))
        legacy_cards = session.scalars(select(StudyCard).where(StudyCard.material_id == 801)).all()
        partial = session.get(Material, 802)
        partial_content = session.scalar(select(MaterialContent).where(MaterialContent.material_id == 802))
        unavailable = session.get(Material, 803)

        assert full is not None and full.extraction_status == "SUCCESS"
        assert full_content is not None and full_content.plain_text == "Memory hierarchy"
        assert full_content.parser_name == "m1-legacy"
        assert full_content.segments[0]["source_reference"] == "旧版提取文本"
        assert legacy_summary is not None and legacy_summary.overview == "Old mock overview"
        assert legacy_summary.model_provider == "legacy-mock"
        assert legacy_summary.prompt_version == "legacy_m1_mock"
        assert len(legacy_cards) == 1
        assert legacy_cards[0].model_provider == "legacy-mock"
        assert partial is not None and partial.extraction_status == "LEGACY_CONTENT_PARTIAL"
        assert partial.extraction_error_code == "LEGACY_TEXT_TRUNCATED"
        assert partial_content is not None and partial_content.parser_name == "m1-legacy-partial"
        assert unavailable is not None and unavailable.extraction_status == "LEGACY_CONTENT_UNAVAILABLE"
        assert unavailable.extraction_error_code == "LEGACY_CONTENT_UNAVAILABLE"
        assert session.scalar(select(MaterialContent).where(MaterialContent.material_id == 803)) is None
    engine.dispose()

    # A second upgrade is a no-op and must not duplicate artifacts.
    migrate_database(tmp_path)
    engine = create_engine(database_url)
    with sessionmaker(bind=engine)() as session:
        assert len(session.scalars(select(MaterialContent)).all()) == 2
        assert len(session.scalars(select(MaterialSummary)).all()) == 1
        assert len(session.scalars(select(StudyCard)).all()) == 1
    engine.dispose()


def test_legacy_import_after_m2_schema_preserves_artifacts_idempotently(tmp_path: Path) -> None:
    source = tmp_path / "legacy-m2.json"
    source.write_text(
        json.dumps(
            {
                "courses": [{"id": 91, "name": "Imported course", "description": ""}],
                "materials": {
                    "91": [
                        {
                            "id": 911,
                            "title": "Saved text",
                            "type": "txt",
                            "summary": "",
                            "extracted": True,
                            "extracted_text": "Pipelining improves throughput.",
                            "text_length": len("Pipelining improves throughput."),
                            "text_preview": "Pipelining improves throughput.",
                            "ai_summary": "Old mock summary",
                            "summary_generated": True,
                            "cards": [{"question": "What improves throughput?", "answer": "Pipelining."}],
                            "cards_generated": True,
                        },
                        {
                            "id": 912,
                            "title": "Preview only",
                            "type": "txt",
                            "summary": "",
                            "extracted": True,
                            "text_preview": "Short preview",
                            "text_length": 200,
                        },
                    ]
                },
            }
        ),
        encoding="utf-8",
    )
    source_before = source.read_bytes()
    database_url = migrate_database(tmp_path)
    options = {
        "source": source,
        "database_url": database_url,
        "dry_run": False,
        "backup_dir": None,
        "bootstrap_user_email": "import-owner@example.test",
        "bootstrap_user_display_name": "Import Owner",
    }
    first = import_legacy_data(**options)
    repeat = import_legacy_data(**options)

    assert first.errors == []
    assert first.courses_created == 1 and first.materials_created == 2
    assert repeat.errors == []
    assert repeat.courses_skipped == 1 and repeat.materials_skipped == 2
    assert source.read_bytes() == source_before

    engine = create_engine(database_url)
    with sessionmaker(bind=engine)() as session:
        saved = session.get(Material, 911)
        preview_only = session.get(Material, 912)
        assert saved is not None and saved.extraction_status == "SUCCESS"
        assert preview_only is not None
        assert preview_only.extraction_status == "LEGACY_CONTENT_UNAVAILABLE"
        assert len(session.scalars(select(MaterialContent)).all()) == 1
        summaries = session.scalars(select(MaterialSummary)).all()
        assert len(summaries) == 1 and summaries[0].model_provider == "legacy-mock"
        cards = session.scalars(select(StudyCard)).all()
        assert len(cards) == 1 and cards[0].model_provider == "legacy-mock"
    engine.dispose()


def test_legacy_import_still_works_before_m2_schema(tmp_path: Path) -> None:
    source = tmp_path / "legacy-m1.json"
    source.write_text(
        json.dumps(
            {
                "courses": [{"id": 71, "name": "M1 course", "description": ""}],
                "materials": {
                    "71": [
                        {
                            "id": 711,
                            "title": "M1 material",
                            "type": "txt",
                            "summary": "",
                            "extracted": True,
                            "text_preview": "Preview",
                            "text_length": 80,
                        }
                    ]
                },
            }
        ),
        encoding="utf-8",
    )
    database_url = migrate_database(tmp_path, "0001_m1a_foundation")
    outcome = import_legacy_data(
        source=source,
        database_url=database_url,
        dry_run=False,
        backup_dir=None,
    )
    assert outcome.errors == []
    assert outcome.courses_created == 1 and outcome.materials_created == 1
    engine = create_engine(database_url)
    assert "material_contents" not in inspect(engine).get_table_names()
    engine.dispose()
