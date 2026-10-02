from __future__ import annotations

import io
import logging
import re
from pathlib import Path
from typing import BinaryIO, Protocol

import pymupdf
from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from backend.app.documents.models import ParsedDocument, ParsedSegment, ParseStatus

logger = logging.getLogger(__name__)
PARSER_VERSION = "m2-v1"


class DocumentParser(Protocol):
    name: str

    def parse(self, payload: bytes) -> ParsedDocument: ...


def _normalize_text(value: str) -> str:
    normalized = value.replace("\r\n", "\n").replace("\r", "\n")
    normalized = re.sub(r"[\t ]+\n", "\n", normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()


def _empty(name: str, page_count: int = 0) -> ParsedDocument:
    return ParsedDocument(
        status=ParseStatus.EMPTY_CONTENT,
        text="",
        segments=[],
        page_count=page_count,
        parser_name=name,
        parser_version=PARSER_VERSION,
        error_code="EMPTY_CONTENT",
        error_message="The document contains no readable text.",
    )


class TextDocumentParser:
    name = "utf8-text"

    def parse(self, payload: bytes) -> ParsedDocument:
        if not payload:
            return _empty(self.name)
        try:
            text = _normalize_text(payload.decode("utf-8"))
        except UnicodeDecodeError:
            logger.warning("Document decoding failed parser=%s", self.name)
            return ParsedDocument(
                status=ParseStatus.DECODE_ERROR,
                text="",
                segments=[],
                page_count=0,
                parser_name=self.name,
                parser_version=PARSER_VERSION,
                error_code="DECODE_ERROR",
                error_message="Uploaded text is not valid UTF-8.",
            )
        if not text:
            return _empty(self.name)
        return ParsedDocument(
            status=ParseStatus.SUCCESS,
            text=text,
            segments=[ParsedSegment(text=text, source_reference="全文")],
            page_count=1,
            parser_name=self.name,
            parser_version=PARSER_VERSION,
        )


class PdfDocumentParser:
    name = "pymupdf"

    def parse(self, payload: bytes) -> ParsedDocument:
        try:
            document = pymupdf.open(stream=payload, filetype="pdf")
        except (pymupdf.FileDataError, RuntimeError, ValueError) as error:
            logger.warning("PDF parsing failed error=%s", type(error).__name__)
            return ParsedDocument(
                status=ParseStatus.PARSE_FAILED,
                text="",
                segments=[],
                page_count=0,
                parser_name=self.name,
                parser_version=PARSER_VERSION,
                error_code="PDF_PARSE_ERROR",
                error_message="The PDF could not be read as a valid document.",
            )
        try:
            segments: list[ParsedSegment] = []
            for page_number, page in enumerate(document, start=1):
                text = _normalize_text(page.get_text("text"))
                if text:
                    segments.append(ParsedSegment(text=text, source_reference=f"PDF 第 {page_number} 页"))
            full_text = "\n\n".join(segment.text for segment in segments)
            if not full_text:
                return _empty(self.name, page_count=len(document))
            return ParsedDocument(
                status=ParseStatus.SUCCESS,
                text=full_text,
                segments=segments,
                page_count=len(document),
                parser_name=self.name,
                parser_version=PARSER_VERSION,
            )
        except (pymupdf.FileDataError, RuntimeError, ValueError) as error:
            logger.warning("PDF extraction failed error=%s", type(error).__name__)
            return ParsedDocument(
                status=ParseStatus.PARSE_FAILED,
                text="",
                segments=[],
                page_count=0,
                parser_name=self.name,
                parser_version=PARSER_VERSION,
                error_code="PDF_PARSE_ERROR",
                error_message="The PDF text could not be extracted.",
            )
        finally:
            document.close()


class DocxDocumentParser:
    name = "python-docx"

    def parse(self, payload: bytes) -> ParsedDocument:
        try:
            document = Document(io.BytesIO(payload))
        except (PackageNotFoundError, ValueError, KeyError, OSError) as error:
            logger.warning("DOCX parsing failed error=%s", type(error).__name__)
            return ParsedDocument(
                status=ParseStatus.PARSE_FAILED,
                text="",
                segments=[],
                page_count=0,
                parser_name=self.name,
                parser_version=PARSER_VERSION,
                error_code="DOCX_PARSE_ERROR",
                error_message="The DOCX file could not be read as a valid document.",
            )

        segments: list[ParsedSegment] = []
        for index, paragraph in enumerate(document.paragraphs, start=1):
            text = _normalize_text(paragraph.text)
            if text:
                segments.append(ParsedSegment(text=text, source_reference=f"DOCX 第 {index} 段"))
        for table_index, table in enumerate(document.tables, start=1):
            rows = []
            for row in table.rows:
                cells = [_normalize_text(cell.text) for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            table_text = _normalize_text("\n".join(rows))
            if table_text:
                segments.append(ParsedSegment(text=table_text, source_reference=f"DOCX 表格 {table_index}"))
        full_text = "\n\n".join(segment.text for segment in segments)
        if not full_text:
            return _empty(self.name)
        return ParsedDocument(
            status=ParseStatus.SUCCESS,
            text=full_text,
            segments=segments,
            page_count=0,
            parser_name=self.name,
            parser_version=PARSER_VERSION,
        )


_PARSERS: dict[str, DocumentParser] = {
    ".txt": TextDocumentParser(),
    ".md": TextDocumentParser(),
    ".pdf": PdfDocumentParser(),
    ".docx": DocxDocumentParser(),
}


def parse_document(source: BinaryIO, extension: str) -> ParsedDocument:
    normalized_extension = Path(f"source{extension}").suffix.lower()
    parser = _PARSERS.get(normalized_extension)
    if parser is None:
        return ParsedDocument(
            status=ParseStatus.UNSUPPORTED_FORMAT,
            text="",
            segments=[],
            page_count=0,
            parser_name="none",
            parser_version=PARSER_VERSION,
            error_code="UNSUPPORTED_FORMAT",
            error_message="Text extraction is not implemented for this file format.",
        )
    return parser.parse(source.read())
